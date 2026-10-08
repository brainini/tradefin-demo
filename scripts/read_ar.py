"""영업팀 미수금 엑셀(d1-4)을 분석용 한 표로 읽는다. 원본은 읽기만 하고 수정하지 않는다.

사용:  uv run python scripts/read_ar.py [엑셀 경로]
       from scripts.read_ar import read_ar;  df = read_ar()
"""
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

DEFAULT_PATH = Path("demos/d1-4_ar_sales_team_2026-09.xlsx")
REGION_SHEETS = ["아시아", "유럽·미주", "중동·아프리카"]  # 요약 시트는 합계만 있어 읽지 않는다
UNIT = 1000  # 원본 금액 단위: 천원 → 원

COLUMNS = ["지역", "담당", "바이어", "국가", "결제조건_원문", "결제조건_코드", "인보이스번호",
           "선적일", "만기일", "선적일_연도추론", "만기일_연도추론", "금액_원", "외화금액", "통화",
           "메모", "원본시트", "원본행", "인보이스_중복"]


def payment_code(raw: str) -> str | None:
    """결제조건 원문 → 표준 코드. 규칙에 없으면 None(확인 필요)."""
    s = re.sub(r"\s+", "", str(raw)).upper()
    if "L/C" in s or s.startswith("LC"):
        return "LC_SIGHT" if "SIGHT" in s else "LC_USANCE" if ("USANCE" in s or "기한부" in s) else None
    if "D/P" in s:
        return "DP"
    if "D/A" in s or s.startswith("DA"):
        return "DA"
    if re.search(r"(TT|T/T)30/70|30%.*70%", s):
        return "TT_SPLIT_30_70"
    if "100%선수" in s or "전액선수" in s:
        return "TT_ADV"
    if s.startswith(("O/A", "OA")) or "사후송금" in s:  # '사후송금'은 O/A로 본다 — 담당자 확인 필요
        return "OA"
    return None


def _to_amount(v) -> float | None:
    """'7,679' 같은 문자열 금액도 숫자로."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and re.fullmatch(r"-?[\d,]+(\.\d+)?", v.strip()):
        return float(v.replace(",", ""))
    return None


def _to_date(v, year_hint: int | None, after: datetime | None = None):
    """날짜형 · '26.8.1' · '8월 7일'. 연도 없는 값은 추론하고 (날짜, 추론여부)로 돌려준다.

    연도 없음: 선적일은 인보이스 번호(INV-YYMM)의 연도, 만기일은 선적일 이후 첫 해당 날짜.
    """
    if isinstance(v, datetime):
        return v, False
    s = str(v).strip() if v is not None else ""
    if m := re.fullmatch(r"(\d\d)\.(\d{1,2})\.(\d{1,2})", s):
        return datetime(2000 + int(m[1]), int(m[2]), int(m[3])), False
    if m := re.fullmatch(r"(\d{1,2})월\s*(\d{1,2})일", s):
        mo, d = int(m[1]), int(m[2])
        if after is not None:
            y = after.year + (1 if (mo, d) < (after.month, after.day) else 0)
        elif year_hint:
            y = year_hint
        else:
            return pd.NaT, False
        return datetime(y, mo, d), True
    return pd.NaT, False


def _find_header_row(ws) -> int:
    for r in range(1, 15):
        if ws.cell(r, 1).value == "담당":
            return r
    raise ValueError(f"{ws.title}: 머리글('담당')을 찾지 못했습니다")


def read_ar(path: str | Path = DEFAULT_PATH) -> pd.DataFrame:
    """네 시트 중 지역 시트 3개를 한 표로. 제목 · 소계 · 합계 · 빈 줄은 건너뛴다.

    2단 머리글(금액 → 원화/외화/통화)은 고정 열 순서로 합쳐 `금액_원`·`외화금액`·`통화`로 쓴다.
    금액은 천원 → 원, 날짜는 datetime. 중복 인보이스는 지우지 않고 `인보이스_중복`만 True로 표시한다.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for name in REGION_SHEETS:
        ws = wb[name]
        header = _find_header_row(ws)  # 머리글은 2단이라 데이터는 header + 2행부터
        for r, v in enumerate(ws.iter_rows(min_row=header + 2, max_col=11, values_only=True), start=header + 2):
            a, inv = v[0], v[4]
            if not a or not inv or "소계" in str(a) or "합계" in str(a):
                continue
            yy = re.fullmatch(r"INV-(\d\d)\d\d-\d+", str(inv))
            hint = 2000 + int(yy[1]) if yy else None
            ship, ship_inf = _to_date(v[5], hint)
            due, due_inf = _to_date(v[6], hint, after=ship if isinstance(ship, datetime) else None)
            won = _to_amount(v[7])
            rows.append({
                "지역": name, "담당": a, "바이어": v[1], "국가": v[2],
                "결제조건_원문": v[3], "결제조건_코드": payment_code(v[3]), "인보이스번호": inv,
                "선적일": ship, "만기일": due, "선적일_연도추론": ship_inf, "만기일_연도추론": due_inf,
                "금액_원": None if won is None else won * UNIT, "외화금액": _to_amount(v[8]), "통화": v[9],
                "메모": v[10], "원본시트": name, "원본행": r,
            })
    wb.close()
    df = pd.DataFrame(rows)
    df["선적일"] = pd.to_datetime(df["선적일"])
    df["만기일"] = pd.to_datetime(df["만기일"])
    df["인보이스_중복"] = df.duplicated("인보이스번호", keep=False)  # 둘 다 표시, 삭제 없음
    return df[COLUMNS]


if __name__ == "__main__":
    df = read_ar(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PATH)
    print(f"읽은 행 수: {len(df)}")
    print(f"중복 표시 건수: {int(df['인보이스_중복'].sum())}행 (인보이스 {df.loc[df['인보이스_중복'], '인보이스번호'].nunique()}건)")
