#!/usr/bin/env python
"""Day4 신용 리스크 리포트 양식(워드 + 같은 내용의 마크다운)을 만들고 검사한다.

만드는 파일(수강생용 = 공개 저장소, 정답 없음)
  labs/day4/d4_risk_report_template.docx  실습 B Step 6–7 양식 — A 설계도(6칸) · B 본문(소제목 8개 + 숫자 대조표) ·
                                          C 레드팀 지적 · D 수정 기록 · E 사람 검토 12항목 + 서명란
  labs/day4/d4_risk_report_template.md    같은 내용의 마크다운판(워드가 없거나 구글 문서 · AI 대화창에 붙여 쓸 때)
손으로 쓴 labs/day4/fraud_signals.md(사기 신호 목록)는 이 스크립트가 만들지 않고 구조만 검사한다.

사용(저장소 루트):
  python tools/build_day4_templates.py           # 생성 + 검사
  python tools/build_day4_templates.py --check   # 생성 없이 검사만
필요 패키지: python-docx. 공개 저장소 파일이다 — 정답 숫자(KPI · 건수 · 금액 · 바이어 ID)를 코드에 적지 않는다.
소제목 · 칸 이름 · 12항목은 labs/day4/prompts.md의 P4-1 [A] · [B]와 같아야 하며, 검사가 대조한다(프롬프트를 고치면 여기도 고친다).
같은 입력이면 같은 파일이 나온다(문서 속성 날짜와 zip 항목 날짜를 고정).
"""
from __future__ import annotations

import argparse
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
LAB4 = REPO / "labs" / "day4"
DOCX = LAB4 / "d4_risk_report_template.docx"
MD = LAB4 / "d4_risk_report_template.md"
FRAUD = LAB4 / "fraud_signals.md"
PROMPTS = LAB4 / "prompts.md"

FONT = "맑은 고딕"
NAVY = "1F3A68"
HEAD_FILL, BOX_FILL = "DDEBF7", "EEF4FB"
FIXED_TIME = datetime(2026, 10, 7, 10, 0, 0)          # 같은 입력이면 같은 파일 — 문서 속성 · zip 항목 날짜 고정
CONTENT_CM = 17.4                                       # A4 폭 21.0 − 좌우 여백 1.8 × 2

# ---------------------------------------------------------------- 내용(워드 · 마크다운 공통)
TITLE = "신용 리스크 요약 리포트 — (가상) 한빛정밀(주)"
LEAD = ("4일차 실습 B Step 6–7 양식 · 저장 이름 d4_credit_report__T조-번호.docx(2조 7번 → d4_credit_report__T2-07.docx). "
        "숫자는 14_ReportInput에서만 가져오고, 문장은 AI 초안을 사람이 고쳐 확정한다.")
GUIDE_TITLE = "작성 안내"
GUIDE = [
    "순서 — A 설계도(Step 6) → B 본문에 P4-1 [A] 초안 붙이기 → C 다른 AI의 P4-1 [B] 지적 붙이기 → 숫자 대조(초록 형광펜) → D 수정 기록 → E 12항목 ✔ → 저장.",
    "숫자는 14_ReportInput에 있는 것만 쓴다. 바이어 실명 · 담당자 개인정보는 쓰지 않는다(buyer_id만).",
    "한도 변경 · 선적 보류 · K-SURE 통지는 '권고(사람 승인 전)'로만 쓴다. 통지 기한은 '결제기일부터 1개월 이내'만 쓰고, 15/1000을 일반 규칙처럼 쓰지 않는다.",
    "[ ]는 채울 자리다. 채운 뒤 지운다. 이 상자는 제출 전에 지워도 된다.",
]
INFO = [
    ("기준일", "[YYYY-MM-DD — 14_ReportInput 표1의 '기준일' 칸 그대로]"),
    ("작성", "[T조-번호]"),
    ("입력", "14_ReportInput 표 5개 — 표1 포트폴리오 KPI · 표2 주의 바이어 Top 5 · 표3 시나리오 · 표4 한도 변경안 · 표5 기한 목록"),
    ("AI 사용", "초안: [AI 서비스 · 기본/추론 모델] · 레드팀: [다른 AI 서비스 · 기본/추론 모델]"),
    ("상태", "AI 작성 초안 → 사람 검토 → 확정 [확정 일시]"),
]
INFO_W = [3.2, 14.2]

# A. 설계도 — 사이트 실습 B Step 6의 표와 같다(앞 세 칸은 채워 두고, 뒤 세 칸은 비운다)
A_TITLE = "A. 리포트 파이프라인 설계도 (Step 6)"
A_GUIDE = ("계산은 엑셀이, 문장은 AI가, 서명은 사람이. 다섯 단계 × 여섯 칸을 채운다 — 실행 방식은 주기 · 조건 · 한 번 중 고르고, "
           "사람 검토 지점은 2곳 이상, 멈춤 규칙은 1개 이상 쓴다.")
A_COLS = ["단계", "입력 파일", "출력 파일", "실행 방식", "사람 검토 지점", "멈춤 규칙"]
A_ROWS = [
    ["① 집계", "d4_end_limits(경보 · 한도 · 시나리오)", "14_ReportInput 표 5개", "", "", ""],
    ["② 재검산", "14_ReportInput", "대조 메모", "", "", ""],
    ["③ 해석(P4-1 [A])", "14_ReportInput", "리포트 초안", "", "", ""],
    ["④ 레드팀(P4-1 [B])", "초안 + 14_ReportInput", "지적 표", "", "", ""],
    ["⑤ 서명", "초안 + 지적 표", "d4_credit_report.docx", "", "", ""],
]
A_W = [2.6, 3.4, 3.0, 2.8, 2.8, 2.8]
A_AFTER = "'위기 바이어 포착 → 경영진 즉시 보고'로 걸려면 어떤 조건에서 실행하나? [한 줄]"

# B. 본문 — P4-1 [A] 출력의 소제목 8개 + 숫자 대조표(labs/day4/prompts.md와 같은 순서 · 이름)
B_TITLE = "B. 리포트 본문 (Step 7 ②)"
B_GUIDE = ("P4-1 [A] 초안을 소제목마다 아래 칸에 옮기고(표는 워드의 '텍스트를 표로 변환') 사람이 고친다. 초안 전체를 붙이고 이 틀을 지워도 된다 — "
           "소제목 8개의 순서와 숫자 대조표는 지킨다. 문장 속 숫자마다 (표-행) 꼬리표를 달고, 입력 표에 없는 것은 '입력에 없음'이라고 쓴다.")
B_SECTIONS = [
    # (소제목, 안내, 종류, 내용) — 종류: lines = 자리 표시 줄, table = (열, 행, 열 너비 cm)
    ("1. 요약 3줄", "가장 중요한 위험 1개 · 권고 조치 1개 · 경영진 결정 요청 1개.", "lines",
     ["① 위험: [한 문장 + (표-행)]", "② 권고 조치: [한 문장 — 권고(사람 승인 전)]", "③ 결정 요청: [한 문장 + 승인 주체]"]),
    ("2. 포트폴리오 현황", "표1에서 옮긴다.", "table",
     (["항목", "값", "입력 표 위치(표-행)"],
      [["총 미결제 채권(USD)"], ["총 미결제 채권(원화)"], ["부보 비중"], ["등급별 바이어 수(S · A · B · C)"], ["C등급 잔액"]],
      [5.4, 5.0, 7.0])),
    ("3. 주의 바이어 Top 5", "표2에서 옮긴다. 사유는 표2의 '상위 사유' 칸 그대로 쓴다.", "table",
     (["buyer_id", "등급", "pd_30d", "미결제액", "경보 단계", "사유 1줄"], [[""]] * 5, [2.4, 1.8, 2.2, 2.8, 2.6, 5.6])),
    ("4. 한도 재배정", "표4에서 옮긴다. 모든 변경은 '권고(사람 승인 전)'다.", "table",
     (["항목", "내용", "입력 표 위치(표-행)"],
      [["현행 한도 합계 → 제안 한도 합계"], ["감액 30% 이상 바이어 수"], ["걸린 제약과 잠재가격 1개(쉬운 말로)"]],
      [5.8, 6.0, 5.6])),
    ("5. 시나리오별 예상손실", "표3에서 옮긴다. 모두 '가정 시나리오에서의 추정'이다. 입력에 없는 칸은 '입력에 없음'.", "table",
     (["시나리오", "EL(USD)", "몬테카를로 평균", "P95", "P99"],
      [["S0"], ["S1"], ["S2"], ["S3"], ["S4"], ["U1(입력에 있으면)"]], [3.4, 3.6, 3.6, 3.4, 3.4])),
    ("6. 조기경보 현황", "표5에서 옮긴다. 통지 여부는 사람이 결정한다. 입력에 없는 칸은 '입력에 없음'.", "table",
     (["구분", "건수", "금액(USD)", "입력 표 위치(표-행)"],
      [["D+7"], ["D+15"], ["D+30(최종 통지 + 선적 보류 — 권고)"], ["선적 보류 바이어 수"],
       ["K-SURE 통지 기한 임박(결제기일부터 1개월 이내)"], ["K-SURE 통지 기한 경과"]],
      [6.4, 2.4, 3.6, 5.0])),
    ("7. 권고 조치와 결정 요청", "조치마다 승인 필요 주체를 쓴다. 상태는 모두 '권고(사람 승인 전)'다.", "table",
     (["조치", "근거(표-행)", "승인 필요 주체", "상태"],
      [["", "", "", "권고(사람 승인 전)"]] * 3, [5.6, 3.6, 4.2, 4.0])),
    ("8. 가정 · 한계 · 출처", "[가정] · [교육용 가정] 항목, 금리 · 환율의 출처와 기준일(입력에 적힌 그대로), 모델 한계 1줄 이상.", "lines",
     ["[가정] · [교육용 가정] 항목: [ ]", "금리 · 환율의 출처와 기준일(입력에 적힌 그대로): [ ]", "모델 한계(1줄 이상): [ ]"]),
]
NUM_TITLE = "숫자 대조표"
NUM_GUIDE = "본문에 쓴 숫자를 빠짐없이 적는다. 입력 값과 맞으면 ✔ 하고 초록 형광펜으로 표시한다. 행이 모자라면 더한다."
NUM_COLS = ["#", "문장 속 숫자", "입력 표 위치(표 번호 - 행 이름)", "입력 값", "대조 ✔"]
NUM_ROWS = [[str(i)] for i in range(1, 7)]
NUM_W = [1.0, 6.2, 5.0, 3.4, 1.8]
B_END = "본 문서는 AI가 작성한 초안이며 담당자 검토 후 확정함"

# C. 레드팀 지적 — P4-1 [B] 출력 표(앞 다섯 칸은 프롬프트와 같고, 마지막 칸은 사람의 판단)
C_TITLE = "C. 레드팀 지적 (Step 7 ③)"
C_GUIDE = ("초안과 다른 AI 서비스에서 P4-1 [B]를 돌려 받은 표를 붙인다. 문제 유형은 7가지 중 하나다 — ① 숫자 불일치 ② 입력에 없는 숫자 "
           "③ 가정 표시 누락 ④ 승인 표기 누락 ⑤ 인과 과장 ⑥ 개인정보 ⑦ 약관 오기. 지적은 후보일 뿐이다 — 입력 표를 직접 열어 맞는 것만 "
           "반영한다. 지적이 없으면 '지적 없음'과 무엇을 어떻게 대조했는지 1줄을 쓴다.")
C_COLS = ["#", "문장 인용(초안 그대로)", "문제 유형", "근거(입력 표 위치 · 값 또는 위 기준)", "고칠 문장 제안", "사람 판단(반영 · 기각 + 이유)"]
C_ROWS = [[str(i)] for i in range(1, 5)]
C_W = [0.9, 4.0, 2.6, 4.0, 3.4, 2.5]
C_END = "확정은 사람이 한다."
RED_TYPES = ["숫자 불일치", "입력에 없는 숫자", "가정 표시 누락", "승인 표기 누락", "인과 과장", "개인정보", "약관 오기"]

# D. 수정 기록 — 제출물 '초안 + 레드팀 지적 + 사람 수정본'(사이트 4.9)을 한 장에서 보이게 한다
D_TITLE = "D. 수정 기록 — AI 초안 → 사람 수정본 (Step 7 ④)"
D_GUIDE = ("사람이 고친 곳을 남긴다 — 제출 파일에는 AI 초안 · 레드팀 지적 · 사람 수정본이 함께 보여야 한다. "
           "이유에는 C의 지적 번호나 E의 항목 번호를 쓴다.")
D_COLS = ["#", "고친 곳(소제목)", "AI 초안 문장", "사람이 고친 문장", "이유(C 번호 · E 번호 · 입력 표 위치)"]
D_ROWS = [[str(i)] for i in range(1, 5)]
D_W = [0.9, 3.0, 5.0, 5.0, 3.5]

# E. 사람 검토 12항목 — labs/day4/prompts.md '사람 검토 체크리스트'와 글자까지 같다(검사가 대조)
E_TITLE = "E. 사람 검토 12항목 (Step 7 ④)"
E_GUIDE = "하나씩 확인하고 ✔ 한다. 1번은 숫자 대조표의 행마다 ✔가 있어야 끝난다. 원문은 과정 사이트 프롬프트 라이브러리 P4-1에 있다."
E_ITEMS = [
    "모든 숫자가 14_ReportInput에 있다(숫자 대조표의 행마다 ✔).",
    "[가정] · [교육용 가정] 표시가 빠지지 않았다.",
    "한도 변경 · 선적 보류 · K-SURE 통지는 \"권고(사람 승인 전)\"로만 썼고 승인 주체가 있다.",
    "K-SURE 사고발생통지 기한 = \"결제기일부터 1개월 이내\", 통지 여부는 사람이 결정한다고 썼다.",
    "15/1000을 일반 규칙처럼 쓰지 않았다(약관 제21조③의 한정 조항).",
    "바이어 실명 · 담당자 개인정보가 없다(buyer_id만).",
    "상관을 원인처럼 쓰지 않았다.",
    "시나리오는 예측이 아니라 \"가정 시나리오에서의 추정\"이라고 썼다.",
    "환율 · 금리의 출처와 기준일을 입력 표에 적힌 대로 썼다.",
    "모델 한계(합성 데이터, 재무 데이터 조작 가능성)를 1줄 이상 적었다.",
    "요약 3줄 · 결론과 표의 숫자가 서로 맞는다.",
    "\"본 문서는 AI가 작성한 초안이며 담당자 검토 후 확정함\" 표시가 있다.",
]
E_COLS = ["번호", "확인 항목", "✔"]
E_W = [1.4, 14.6, 1.4]
SIGN_COLS = ["", "작성", "검토(담당자)", "승인(승인 주체)"]
SIGN_ROWS = [["이름", "[T조-번호]", "", ""], ["일시", "", "", ""]]
SIGN_W = [2.4, 5.0, 5.0, 5.0]
FOOT = "사람 승인 전 초안 — 한도 변경 · 선적 보류 · K-SURE 통지는 승인 주체가 확정한다."
HEADER_TXT = "신용 리스크 요약 리포트 · 4일차 양식 · AI 작성 초안(사람 승인 전)"


def _pad(row, n):
    return list(row) + [""] * (n - len(row))


# ---------------------------------------------------------------- 마크다운
def _md_table(cols, rows):
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        out.append("| " + " | ".join(c if c else " " for c in _pad(r, len(cols))) + " |")
    return out


def build_md() -> Path:
    o = [f"# {TITLE}", "", f"> {LEAD}", "", f"> **{GUIDE_TITLE}**", ">"]
    o += [f"> - {g}" for g in GUIDE]
    o += ["", "| 항목 | 내용 |", "|---|---|"] + [f"| {k} | {v} |" for k, v in INFO]
    o += ["", f"## {A_TITLE}", "", f"*{A_GUIDE}*", ""] + _md_table(A_COLS, A_ROWS) + ["", A_AFTER]
    o += ["", f"## {B_TITLE}", "", f"*{B_GUIDE}*"]
    for head, guide, kind, body in B_SECTIONS:
        o += ["", f"### {head}", "", f"*{guide}*", ""]
        if kind == "lines":
            o += [f"- {t}" for t in body]
        else:
            cols, rows, _w = body
            o += _md_table(cols, rows)
            if head.startswith("5."):
                o += ["", "원화 영향 1줄: [ ]"]
    o += ["", f"### {NUM_TITLE}", "", f"*{NUM_GUIDE}*", ""] + _md_table(NUM_COLS, NUM_ROWS)
    o += ["", f"**{B_END}**"]
    o += ["", f"## {C_TITLE}", "", f"*{C_GUIDE}*", ""] + _md_table(C_COLS, C_ROWS) + ["", C_END]
    o += ["", f"## {D_TITLE}", "", f"*{D_GUIDE}*", ""] + _md_table(D_COLS, D_ROWS)
    o += ["", f"## {E_TITLE}", "", f"*{E_GUIDE}*", ""]
    o += _md_table(E_COLS, [[str(i), t] for i, t in enumerate(E_ITEMS, 1)])
    o += ["", "---", ""] + _md_table(SIGN_COLS, SIGN_ROWS) + ["", FOOT, ""]
    MD.write_text("\n".join(o), encoding="utf-8")
    return MD


# ---------------------------------------------------------------- 워드
def _fix_style_font(style) -> None:
    from docx.oxml.ns import qn

    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):   # 테마 글꼴이 남으면 Word가 그쪽을 쓴다
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _write(cell, text, bold=False, gray=False, size=None) -> None:
    from docx.shared import Pt, RGBColor

    p = cell.paragraphs[0]
    p.style = "Table Text"                      # 빈 칸에 글을 써도 같은 글자 크기가 이어지게 문단 스타일로 지정
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if not text:
        return
    run = p.add_run(text)
    if bold:
        run.bold = True
    if gray:
        run.italic = True
        run.font.color.rgb = RGBColor(0x76, 0x76, 0x76)
    if size:
        run.font.size = Pt(size)


def _grid(doc, cols, rows, widths, min_row_cm=0.9, head_fill=HEAD_FILL):
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    assert abs(sum(widths) - CONTENT_CM) < 0.05 and len(widths) == len(cols), (cols, widths)
    t = doc.add_table(rows=1 + len(rows), cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, w in enumerate(widths):
        t.columns[j].width = Cm(w)
    for i, row in enumerate([cols] + [_pad(r, len(cols)) for r in rows]):
        tr = t.rows[i]
        trpr = tr._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trpr.append(hdr)
        else:
            tr.height, tr.height_rule = Cm(min_row_cm), WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, text in enumerate(row):
            c = tr.cells[j]
            c.width = Cm(widths[j])
            if i == 0:
                _write(c, text, bold=True)
                _shade(c, head_fill)
            else:
                _write(c, text)
    return t


def _field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    parts = []
    for kind in ("begin", None, "separate", "t", "end"):
        if kind is None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        elif kind == "t":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        parts.append(el)
    for el in parts:
        run._r.append(el)


def _normalize_zip(path: Path) -> None:
    """같은 내용이면 같은 바이트 — zip 항목 날짜를 고정하고 이름 순으로 다시 쓴다([Content_Types].xml을 맨 앞에)."""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i)) for i in z.infolist()]
    items.sort(key=lambda x: (x[0] != "[Content_Types].xml", x[0]))
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, data)
    tmp.replace(path)


def build_docx() -> Path:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
    sec.header_distance = sec.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    _fix_style_font(normal)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    cell_style = doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    cell_style.base_style = normal
    cell_style.font.size = Pt(9.5)
    cell_style.paragraph_format.space_after = Pt(0)
    for name, size, before in (("Heading 1", 16, 0), ("Heading 2", 13, 10), ("Heading 3", 11, 8)):
        s = doc.styles[name]
        _fix_style_font(s)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(NAVY)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True

    def gray(text, after=4):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x76, 0x76, 0x76)
        p.paragraph_format.space_after = Pt(after)
        return p

    def h2(text, new_page=True):
        p = doc.add_heading(text, level=2)
        p.paragraph_format.page_break_before = new_page
        return p

    # 머리글 · 바닥글(AI 초안 표시와 쪽 번호)
    hp = sec.header.paragraphs[0]
    hr_ = hp.add_run(HEADER_TXT)
    hr_.font.size = Pt(8.5)
    hr_.font.color.rgb = RGBColor(0x76, 0x76, 0x76)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("쪽 ")
    fr.font.size = Pt(8.5)
    _field(fp, "PAGE")
    for r in fp.runs:
        r.font.size = Pt(8.5)

    # 표지 부분: 제목 · 안내 상자 · 기본 정보
    doc.add_heading(TITLE, level=1)
    gray(LEAD, after=6)
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.autofit = False
    box.columns[0].width = Cm(CONTENT_CM)
    cell = box.cell(0, 0)
    cell.width = Cm(CONTENT_CM)
    _shade(cell, BOX_FILL)
    _write(cell, GUIDE_TITLE, bold=True, size=10)
    for line in GUIDE:
        p = cell.add_paragraph()
        r = p.add_run("· " + line)
        r.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    info = _grid(doc, ["항목", "내용"], [[k, v] for k, v in INFO], INFO_W, min_row_cm=0.8)
    for row in info.rows[1:]:
        for r in row.cells[0].paragraphs[0].runs:
            r.bold = True

    # A 설계도
    h2(A_TITLE)
    gray(A_GUIDE)
    _grid(doc, A_COLS, A_ROWS, A_W, min_row_cm=1.6)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run(A_AFTER)

    # B 본문
    h2(B_TITLE)
    gray(B_GUIDE)
    for head, guide, kind, body in B_SECTIONS:
        doc.add_heading(head, level=3)
        gray(guide, after=3)
        if kind == "lines":
            for t in body:
                doc.add_paragraph(t)
        else:
            cols, rows, widths = body
            _grid(doc, cols, rows, widths, min_row_cm=0.8)
            if head.startswith("5."):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.add_run("원화 영향 1줄: [ ]")
    doc.add_heading(NUM_TITLE, level=3)
    gray(NUM_GUIDE, after=3)
    _grid(doc, NUM_COLS, NUM_ROWS, NUM_W, min_row_cm=0.8)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run(B_END).bold = True

    # C 레드팀 · D 수정 기록
    h2(C_TITLE)
    gray(C_GUIDE)
    _grid(doc, C_COLS, C_ROWS, C_W, min_row_cm=1.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.add_run(C_END)
    h2(D_TITLE, new_page=False)
    gray(D_GUIDE)
    _grid(doc, D_COLS, D_ROWS, D_W, min_row_cm=1.5)

    # E 12항목 · 서명
    h2(E_TITLE)
    gray(E_GUIDE)
    _grid(doc, E_COLS, [[str(i), t] for i, t in enumerate(E_ITEMS, 1)], E_W, min_row_cm=0.8)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(10)
    _grid(doc, SIGN_COLS, SIGN_ROWS, SIGN_W, min_row_cm=1.0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run(FOOT).bold = True

    cp = doc.core_properties
    cp.title = TITLE
    cp.author = "tradefin-ai-2026"
    cp.last_modified_by = "tradefin-ai-2026"
    cp.comments = "4일차 실습 B Step 6–7 양식 — 설계도 6칸 · 본문 8소제목 · 레드팀 지적 · 수정 기록 · 12항목"
    cp.created = cp.modified = FIXED_TIME
    cp.revision = 1
    doc.save(DOCX)
    _normalize_zip(DOCX)
    return DOCX


# ---------------------------------------------------------------- 검사
def _between(text: str, key: str) -> str:
    m = re.search(rf"<!-- --8<-- \[start:{re.escape(key)}\] -->(.*?)<!-- --8<-- \[end:{re.escape(key)}\] -->", text, re.S)
    return m.group(1) if m else ""


def check() -> list[str]:
    from docx import Document

    probs: list[str] = []
    md = MD.read_text(encoding="utf-8")
    d = Document(str(DOCX))
    h2 = [p.text for p in d.paragraphs if p.style.name == "Heading 2"]
    h3 = [p.text for p in d.paragraphs if p.style.name == "Heading 3"]
    want2 = [A_TITLE, B_TITLE, C_TITLE, D_TITLE, E_TITLE]
    want3 = [s[0] for s in B_SECTIONS] + [NUM_TITLE]
    if h2 != want2 or h3 != want3:
        probs.append(f"워드 소제목 불일치: {h2} {h3}")
    for h in want2:
        if f"\n## {h}\n" not in md:
            probs.append(f"마크다운에 없는 제목: {h}")
    for h in want3:
        if f"\n### {h}\n" not in md:
            probs.append(f"마크다운에 없는 소제목: {h}")
    n_tables = 1 + 1 + 1 + sum(1 for s in B_SECTIONS if s[2] == "table") + 1 + 1 + 1 + 1 + 1   # 상자 · 정보 · A · B표 · 숫자 · C · D · E · 서명
    if len(d.tables) != n_tables:
        probs.append(f"워드 표 개수 {len(d.tables)} ≠ {n_tables}")
    else:
        a = d.tables[2]
        if [c.text for c in a.rows[0].cells] != A_COLS or len(a.rows) != 1 + len(A_ROWS):
            probs.append("설계도 표가 6칸 × 5단계가 아님")
        e = d.tables[-2]
        if [r.cells[1].text for r in e.rows[1:]] != E_ITEMS:
            probs.append("워드 12항목 글자가 E_ITEMS와 다름")
    for t in (A_ROWS, NUM_ROWS, C_ROWS, D_ROWS):
        for r in t:
            if any("|" in c for c in r):
                probs.append(f"표 칸에 '|'가 있음: {r}")

    # 프롬프트(labs/day4/prompts.md)와 같은가
    p = PROMPTS.read_text(encoding="utf-8")
    blk_a, blk_b = _between(p, "p4-1a"), _between(p, "p4-1b")
    heads = [re.sub(r"\s+—.*$", "", ln[3:]).strip() for ln in blk_a.splitlines() if ln.startswith("## ")]
    if heads != want3:
        probs.append(f"P4-1 [A] 소제목과 양식 소제목이 다름: {heads}")
    num_line = next((ln for ln in blk_a.splitlines() if ln.startswith("## 숫자 대조표")), "")
    for c in NUM_COLS[:4]:
        if c not in num_line:
            probs.append(f"P4-1 [A] 숫자 대조표 열 이름과 다름: {c}")
    hdr_b = next((ln for ln in blk_b.splitlines() if ln.startswith("| # |")), "")
    if [c.strip() for c in hdr_b.strip("|").split("|")] != C_COLS[:5]:
        probs.append(f"P4-1 [B] 지적 표 머리와 다름: {hdr_b}")
    for t in RED_TYPES:
        if t not in blk_b:
            probs.append(f"P4-1 [B]에 없는 문제 유형: {t}")
    items = [re.sub(r"`", "", x) for x in re.findall(r"^- \[ \] \d+\. (.+)$", p, flags=re.M)]
    if items != E_ITEMS:
        probs.append(f"사람 검토 12항목이 prompts.md와 다름({len(items)}개)")
    for key in ("권고(사람 승인 전)", "결제기일부터 1개월 이내", B_END):
        if key not in md or key not in p:
            probs.append(f"고정 문구 누락: {key}")

    # 공개 저장소 규칙 — 정답 숫자 · 바이어 ID · 실명 없음
    src = Path(__file__).read_text(encoding="utf-8")
    for label, text in (("마크다운", md), ("스크립트", src), ("워드", "\n".join(p_.text for p_ in d.paragraphs))):
        if re.search(r"\b\d{1,3}(,\d{3})+\b", text) or re.search(r"\bB\d{3}\b", text):
            probs.append(f"{label}에 정답 숫자 · 바이어 ID 형태가 있음")
    return probs + _check_fraud()


def _check_fraud() -> list[str]:
    """fraud_signals.md — 점수 열 · 할 일 열 · 조치 사다리가 있고, 바이어 ID · 위기 카드 내용이 없는가."""
    if not FRAUD.is_file():
        return [f"없음: {FRAUD.relative_to(REPO)}"]
    t = FRAUD.read_text(encoding="utf-8")
    probs = []
    for key in ("점수", "할 일", "멈춤", "확인", "기록", "권고", "[교육용 가정]"):
        if key not in t:
            probs.append(f"fraud_signals.md에 '{key}'가 없음")
    if re.search(r"\bB\d{3}\b|NEW\d|카드\s?[①-⑤1-5]", t):
        probs.append("fraud_signals.md에 바이어 ID · 위기 카드 형태가 있음")
    return probs


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--check", action="store_true", help="생성하지 않고 검사만")
    a = ap.parse_args()
    if not a.check:
        print("썼음:", build_md().relative_to(REPO))
        print("썼음:", build_docx().relative_to(REPO))
    probs = check()
    if probs:
        print("검사 실패:", *probs, sep="\n  - ")
        return 1
    print("검사 통과 — 설계도 6칸 × 5단계 · 본문 소제목 8 + 숫자 대조표 = P4-1 [A] · 지적 표 = P4-1 [B] · 12항목 = prompts.md · 워드 = 마크다운")
    return 0


if __name__ == "__main__":
    sys.exit(main())
