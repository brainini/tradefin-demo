#!/usr/bin/env python
"""Day3 Orange 완성 워크플로 `labs/day3/d3_orange_model.ows`를 만들고, 수강생 PC처럼 다시 열어 끝까지 실행해 검사한다.

정본: DAY3_SPEC_v3 부록 B Step 1–7 · 03 Part2 §3.5(위젯 체인·설정) · 사이트 docs/day3/lab-a.md '워크플로 한 장'.
  [File] d3_start_features.csv (late_30d = target · buyer_id · ref_date · sample_weight · split_suggested = meta ·
         fx_krw_close · usd_krw_close = skip)
   └ [Data Sampler ①] 70% · Stratify · Replicable ─ Data Sample ─┬─ [Select Rows] late_30d is 1 ─ [Data Sampler ②]
                                                                │    고정 개수 = 학습 연체 × 2 · 복원 추출 ─┐
                                                                └──────────────────────── [Concatenate] ←─┘ = 학습셋
                      └ Remaining Data(시험 30%, 늘리지 않음) → [Test and Score] Test Data · [Predictions (시험셋)]
  [Gradient Boosting] xgboost 100 · 0.1 · 깊이 3 · Replicable [교육용 가정] · [Logistic Regression] 기준선
   → [Test and Score] Test on test data · Target 1 → [Confusion Matrix] → [Data Table (FN)] · [ROC Analysis] FP 1 · FN 5
     · [Calibration Plot] Target 1 · Sigmoid · 모델 1개(Gradient Boosting) → 출력 모델 → [Predictions] ×2
  [File ②] d3_start_scoring.csv → [Predictions] → [Save Data] d3_predictions.xlsx · [Data Table (예측)] → [Explain Prediction]
  [Gradient Boosting] Model → [Explain Model](Data = 학습 70%, Target 1, 상위 10) · [Explain Prediction](Background = 학습 70%)

사용(저장소 루트, Orange 3.40 환경 — tools/orange_ows.py 머리말의 패키지):
  python tools/build_day3_orange.py              # .ows 생성 → 다시 열기 → 수강생 PC처럼 실행 검사(요약 출력)
  python tools/build_day3_orange.py --expected   # + 강사 리허설 기대값 ../instructor/day3/answers/orange_expected.json
공개 저장소 파일이다 — 정답 숫자를 코드에 적지 않는다(기대값은 실행해서 강사 폴더에만 쓴다).
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
import orange_ows as ow  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
CK = REPO / "data" / "checkpoints"
OUT = REPO / "labs" / "day3" / "d3_orange_model.ows"
ANSWERS = REPO.parent / "instructor" / "day3" / "answers"
TRAIN = "d3_start_features.csv"
SCORE = "d3_start_scoring.csv"

ROLES_TRAIN = {"late_30d": "target", "buyer_id": "meta", "ref_date": "meta", "sample_weight": "meta",
               "split_suggested": "meta", "fx_krw_close": "skip", "usd_krw_close": "skip"}
ROLES_SCORE = {"buyer_id": "meta", "ref_date": "meta", "fx_krw_close": "skip", "usd_krw_close": "skip"}
GB = "Orange.widgets.model.owgradientboosting.OWGradientBoosting"
Q = {
    "file": "Orange.widgets.data.owfile.OWFile",
    "table": "Orange.widgets.data.owtable.OWTable",
    "sampler": "Orange.widgets.data.owdatasampler.OWDataSampler",
    "select": "Orange.widgets.data.owselectrows.OWSelectRows",
    "concat": "Orange.widgets.data.owconcatenate.OWConcatenate",
    "gb": GB,
    "lr": "Orange.widgets.model.owlogisticregression.OWLogisticRegression",
    "ts": "Orange.widgets.evaluate.owtestandscore.OWTestAndScore",
    "cm": "Orange.widgets.evaluate.owconfusionmatrix.OWConfusionMatrix",
    "roc": "Orange.widgets.evaluate.owrocanalysis.OWROCAnalysis",
    "cal": "Orange.widgets.evaluate.owcalibrationplot.OWCalibrationPlot",
    "pred": "Orange.widgets.evaluate.owpredictions.OWPredictions",
    "save": "Orange.widgets.data.owsave.OWSave",
    "xm": "orangecontrib.explain.widgets.owexplainmodel.OWExplainModel",
    "xp": "orangecontrib.explain.widgets.owexplainprediction.OWExplainPrediction",
}
T = {  # 노드 제목 = 사이트·Step 카드의 이름
    "file": "File", "ds1": "Data Sampler ①", "t_test": "Data Table (시험셋)", "sel": "Select Rows",
    "t_pos": "Data Table (학습 연체)", "ds2": "Data Sampler ②", "concat": "Concatenate",
    "t_train": "Data Table (학습셋 오버샘플)", "gb": "Gradient Boosting", "lr": "Logistic Regression",
    "ts": "Test and Score", "cm": "Confusion Matrix", "t_fn": "Data Table (FN)", "roc": "ROC Analysis",
    "cal": "Calibration Plot", "file2": "File ②", "pred": "Predictions", "save": "Save Data",
    "t_pred": "Data Table (예측)", "pred_t": "Predictions (시험셋)", "save_t": "Save Data (시험셋)",
    "xm": "Explain Model", "xp": "Explain Prediction",
}


NOTES = [
    "Orange 3.40 Calibration Plot은 '보정 모델' 또는 '임계값 모델' 하나만 낸다: 그래프가 Calibration curve면 Sigmoid 보정 모델(임계값 0.5), "
    "지표 그래프(예: Sensitivity and specificity)에서 세로선을 끌면 보정하지 않은 원래 모델 + 임계값 t(ThresholdClassifier). "
    "그래서 Predictions의 확률은 마지막에 보던 그래프에 따라 달라진다 — 등급 매칭에는 같은 척도에서 읽은 t를 쓴다.",
    "Predictions는 라벨이 없는 점수화 파일(d3_start_scoring.csv)에 기본값으로 확률 열을 내지 않는다 → .ows에 'Classes known to the model'로 저장해 두었다(열 이름 끝 '(1)').",
    "Save Data(xlsx) 첫 열은 Predictions가 붙인 'Selected', 그다음 ref_date · buyer_id · 예측 열 순서다.",
    "xgboost 3.3부터 sklearn 기본값 enable_categorical=True → shap(0.50+)이 Explain Model · Explain Prediction을 거부('Categorical split is not yet supported'). "
    "Orange는 범주를 원-핫으로 바꿔 넣으므로 모델은 같다 — 포터블 Orange의 xgboost를 3.2.0 이하로 두거나, 그 자리에서 Gradient Boosting을 scikit-learn 방식(같은 100·0.1·3)으로 바꾼다.",
    "xgboost가 없으면 Gradient Boosting 위젯은 조용히 scikit-learn 방식으로 바뀐다(설정 100 · 0.1 · 3은 같은 값으로 저장해 두었다).",
    "Explain Prediction의 '예측'은 보정 전 모델의 확률이라 pd_30d(보정 모델)와 다를 수 있다 — 해석서에는 기여 순서와 부호를 쓴다.",
]


def xlsx_filter() -> str:
    from Orange.widgets.data.owsave import OWSave

    return next(f for f in OWSave.get_filters() if f.endswith("(*.xlsx)") and not f.startswith("Compressed"))


def build(workdir: Path) -> tuple[Path, dict]:
    """임시 D3 폴더(basedir)에서 캔버스를 만들고 .ows를 쓴다. 만들 때 본 값(학습 연체 수 등)을 돌려준다."""
    d3 = workdir / "D3"
    d3.mkdir(parents=True, exist_ok=True)
    for name in (TRAIN, SCORE):
        shutil.copy2(CK / name, d3 / name)
    fl = ow.Flow("3일차 — 30일 이내 연체 확률 모델(Orange 3.40)",
                 "DAY3 블록 A Step 1–4 · 블록 B Step 5–7. 과정 합성 데이터((가상) 한빛정밀 바이어 150곳). "
                 "File 위젯을 D3 폴더의 d3_start_features.csv로 다시 지정하면 전체가 다시 계산된다. "
                 "파라미터(70/30 · 오버샘플 ×2 · xgboost 100·0.1·3 · FN:FP = 5:1)는 [교육용 가정].", d3)
    seen: dict = {}

    # ---- Step 1: File → 역할 → Data Sampler ①
    f = fl.add("file", Q["file"], T["file"], (0, 260), {"recent_paths": [ow.recent_path(TRAIN)], "source": 0})
    fl.settle()
    ow.set_roles(f, ROLES_TRAIN, {"late_30d": "categorical"})
    fl.settle()
    fl.add("ds1", Q["sampler"], T["ds1"], (170, 260),
           {"sampling_type": 0, "sampleSizePercentage": 70, "stratify": True, "use_seed": True, "replacement": False})
    fl.link("file", "data", "ds1", "data")
    fl.add("t_test", Q["table"], T["t_test"], (340, 400))
    fl.link("ds1", "remaining_data", "t_test", "data")
    fl.settle()

    # ---- Step 2: Select Rows(late_30d is 1) → Data Sampler ②(×2, 복원) → Concatenate
    sel = fl.add("sel", Q["select"], T["sel"], (340, 130), {"purge_attributes": False, "purge_classes": False})
    fl.link("ds1", "data_sample", "sel", "data")
    fl.settle()
    y = sel.data.domain.class_var
    sel.add_row(y, 0, [y.values.index("1") + 1])          # 'is' · 값 '1'(위젯 안 번호는 1부터)
    sel.conditions_changed()
    sel.commit.now()
    fl.add("t_pos", Q["table"], T["t_pos"], (510, 30))
    fl.link("sel", "matching_data", "t_pos", "data")
    fl.settle()
    pos = ow.link_value(fl.scheme, T["sel"], "matching_data")
    seen["train_pos"] = int(len(pos))
    n_over = 2 * seen["train_pos"]                         # 학습 연체 × 2 → 연체가 3배 [교육용 가정]
    fl.add("ds2", Q["sampler"], T["ds2"], (510, 130),
           {"sampling_type": 1, "sampleSizeNumber": n_over, "replacement": True, "use_seed": True, "stratify": False})
    fl.link("sel", "matching_data", "ds2", "data")
    fl.add("concat", Q["concat"], T["concat"], (680, 260))
    fl.link("ds1", "data_sample", "concat", "primary_data")
    fl.link("ds2", "data_sample", "concat", "additional_data")
    fl.add("t_train", Q["table"], T["t_train"], (850, 130))
    fl.link("concat", "data", "t_train", "data")
    fl.settle()

    # ---- Step 3: 모델 · Test and Score(Test on test data · Target 1) · Confusion Matrix → Data Table(FN)
    gb_defaults = {"colsample_bylevel": 1, "colsample_bynode": 1, "colsample_bytree": 1, "lambda_index": 53,
                   "subsample": 1}
    fl.add("gb", Q["gb"], T["gb"], (850, 260),
           {"method_index": 1, "auto_apply": True,
            "xgb_editor": {**gb_defaults, "n_estimators": 100, "learning_rate": 0.1, "max_depth": 3, "random_state": True},
            "gb_editor": {"n_estimators": 100, "learning_rate": 0.1, "max_depth": 3, "min_samples_split": 2,
                          "random_state": True, "subsample": 1}})
    fl.link("concat", "data", "gb", "data")
    fl.add("lr", Q["lr"], T["lr"], (850, 400))
    fl.add("ts", Q["ts"], T["ts"], (1020, 300), {"resampling": 5})   # 5 = Test on test data
    fl.link("concat", "data", "ts", "train_data")
    fl.link("ds1", "remaining_data", "ts", "test_data")
    fl.link("gb", "learner", "ts", "learner")                        # 학습기 0 = Gradient Boosting
    fl.link("lr", "learner", "ts", "learner")                        # 학습기 1 = Logistic Regression(기준선)
    fl.settle()
    ts = fl.w("ts")
    ts.class_selection = "1"
    ts._on_target_class_changed()
    fl.add("cm", Q["cm"], T["cm"], (1190, 170), {"selected_learner": [0], "selected_quantity": 0})
    fl.link("ts", "evaluations_results", "cm", "evaluation_results")
    fl.add("t_fn", Q["table"], T["t_fn"], (1360, 170))
    fl.link("cm", "selected_data", "t_fn", "data")

    # ---- Step 4: ROC Analysis(FP 1 · FN 5) · Calibration Plot(Target 1 · Sigmoid · 모델 1개)
    fl.add("roc", Q["roc"], T["roc"], (1190, 300), {"fp_cost": 1, "fn_cost": 5, "display_perf_line": True,
                                                    "display_def_threshold": True})
    fl.link("ts", "evaluations_results", "roc", "evaluation_results")
    fl.add("cal", Q["cal"], T["cal"], (1190, 430), {"output_calibration": 0, "score": 0, "threshold": 0.5,
                                                    "display_rug": True})
    fl.link("ts", "evaluations_results", "cal", "evaluation_results")
    fl.settle()
    if ow.link_value(fl.scheme, T["ts"], "evaluations_results") is None:
        raise RuntimeError("Test and Score가 결과를 내지 않았다 — 위 경고(Test and Score)를 본다")
    roc = fl.w("roc")
    roc.target_index = 1
    roc._on_target_changed()
    cal = fl.w("cal")
    cal.target_index = 1
    cal.target_index_changed()
    cal.selected_classifiers = [0]
    cal._on_selection_changed()
    fl.settle()

    # ---- Step 5: File ② → Predictions(보정 모델) → Save Data · Data Table(예측)
    f2 = fl.add("file2", Q["file"], T["file2"], (1190, 620), {"recent_paths": [ow.recent_path(SCORE)], "source": 0})
    fl.settle()
    ow.set_roles(f2, ROLES_SCORE)
    fl.add("pred", Q["pred"], T["pred"], (1360, 560))
    fl.link("file2", "data", "pred", "data")
    fl.link("cal", "calibrated_model", "pred", "predictors")
    xf = xlsx_filter()
    fl.add("save", Q["save"], T["save"], (1530, 520),
           {"filter": xf, "stored_name": "d3_predictions.xlsx", "stored_path": ".", "auto_save": False,
            "add_type_annotations": False})
    fl.link("pred", "annotated", "save", "data")
    fl.add("t_pred", Q["table"], T["t_pred"], (1530, 640))
    fl.link("pred", "annotated", "t_pred", "data")
    fl.settle()
    # 점수화 파일에는 라벨이 없어 기본값('Classes in data')으로는 확률 열이 하나도 안 나온다 → 'Classes known to the model'
    pw = fl.w("pred")
    pw.controls.shown_probs.setCurrentIndex(pw.MODEL_PROBS)
    pw.controls.shown_probs.activated.emit(pw.MODEL_PROBS)
    fl.settle()
    # ---- Step 6 ④: 시험셋 확률(단조성 점검용) — 사이트 lab-b Step 6 ④
    fl.add("pred_t", Q["pred"], T["pred_t"], (1360, 430))
    fl.link("ds1", "remaining_data", "pred_t", "data")
    fl.link("cal", "calibrated_model", "pred_t", "predictors")
    fl.add("save_t", Q["save"], T["save_t"], (1530, 400),
           {"filter": xf, "stored_name": "d3_test_pred.xlsx", "stored_path": ".", "auto_save": False,
            "add_type_annotations": False})
    fl.link("pred_t", "annotated", "save_t", "data")

    # ---- Step 7: Explain Model · Explain Prediction
    fl.add("xm", Q["xm"], T["xm"], (1020, 130), {"n_attributes": 10, "show_legend": True})
    fl.link("ds1", "data_sample", "xm", "data")
    fl.link("gb", "model", "xm", "model")
    fl.add("xp", Q["xp"], T["xp"], (1700, 640))
    fl.link("gb", "model", "xp", "model")
    fl.link("ds1", "data_sample", "xp", "background_data")
    fl.link("t_pred", "selected_data", "xp", "data")
    fl.settle(timeout=600)
    xm = fl.w("xm")
    xm.controls.target_index.setCurrentIndex(1)
    xm.controls.target_index.activated.emit(1)
    fl.settle(timeout=600)
    # Explain Prediction의 Target도 1로 — 행이 들어와야 문맥이 열리므로 한 행을 잠깐 고른 뒤 설정하고 선택은 다시 비운다
    tp = fl.w("t_pred")
    tbl = ow.link_value(fl.scheme, T["pred"], "annotated")
    ncol = len(tbl.domain.variables) + len(tbl.domain.metas)
    tp.set_selection([0], list(range(ncol)))
    tp.commit.now()
    fl.settle(timeout=600)
    xp = fl.w("xp")
    xp.controls.target_index.setCurrentIndex(1)
    xp.controls.target_index.activated.emit(1)
    fl.settle(timeout=600)
    tp.clear_selection()
    tp.commit.now()
    fl.settle()

    # ---- 캔버스 메모(수강생이 여는 순서대로)
    fl.note((-20, -150, 700, 60),
            "3일차 30일 연체 확률 모델 — 블록 A Step 1–4 · 블록 B Step 5–7. 처음 열면 File 두 개를 D3 폴더의 CSV로 "
            "다시 지정합니다(경고 ! 는 그 뒤 사라짐). 숫자 설정은 [교육용 가정] — 오늘은 바꾸지 않습니다.", 13)
    fl.note((-40, 330, 170, 190),
            "Step 1 · File 더블클릭 → 폴더 아이콘 → d3_start_features.csv\n역할: late_30d = target · buyer_id · "
            "ref_date · sample_weight · split_suggested = meta · fx_krw_close · usd_krw_close = skip → Apply")
    fl.note((140, 470, 200, 80), "Data Sampler ① · 70% · Stratify ✔ · Replicable ✔\nRemaining Data(시험 30%)는 "
                                  "어떤 위젯으로도 늘리지 않는다")
    fl.note((330, -70, 380, 60), "Step 2 · 학습셋 연체만 복제: Select Rows(late_30d is 1) → Data Sampler ② "
                                 f"고정 개수 {n_over}(= 학습 연체 × 2) · 복원 추출 → Concatenate")
    fl.note((740, 470, 240, 80), "Step 3 · Gradient Boosting = xgboost · 나무 100 · 학습률 0.1 · 깊이 3 · "
                                  "Replicable. xgboost가 회색이면 scikit-learn 방식(같은 값)으로 바뀐다")
    fl.note((940, 370, 160, 60), "Test on test data · Target class 1\n(교차검증이면 임계값 모델이 안 나온다)")
    fl.note((1250, 270, 230, 50), "Step 4 · ROC: Target 1 · FP 비용 1 · FN 비용 5 → 성능선이 닿는 점")
    fl.note((1020, 490, 230, 100), "Calibration Plot · Sigmoid로 곡선 확인 → 지표를 Sensitivity and specificity로 "
                                  "바꿔 세로선 = t → 출력 모델이 Predictions로")
    fl.note((1330, 710, 420, 50), "Step 5 · File ② = d3_start_scoring.csv(150행) → Predictions → Save Data "
                                  "d3_predictions.xlsx · Step 7 · Data Table에서 1행 선택 → Explain Prediction")
    fl.note((1600, 360, 220, 50), "Step 6 ④ · 시험셋 확률 d3_test_pred.xlsx — 단조성 점검용")
    fl.note((970, -40, 280, 100), "Step 7 · Explain Model · Target 1 · 상위 10개 특성(학습 70% 기준). "
                                 "'Categorical split' 오류가 나면 강사 안내 — Gradient Boosting을 scikit-learn 방식"
                                 "(같은 값)으로 바꾸면 설명이 나온다")

    # Select Rows: 데이터가 다시 들어올 때마다 문맥에서 행을 되살리므로, 저장 직전에 조건을 '하나'로 다시 맞춘다
    sel.remove_all_rows()
    sel.add_row(y, 0, [y.values.index("1") + 1])
    sel.conditions_changed()
    sel.commit.now()
    fl.settle()
    if len(sel.conditions) != 1:
        raise RuntimeError(f"Select Rows 조건이 {len(sel.conditions)}개: {sel.conditions}")

    seen.update({"n_over": n_over, "xlsx_filter": xf})
    fl.save(OUT)
    fl.close()
    return OUT, seen


# ---------------------------------------------------------------- 검사 · 리허설 기대값
def confusion(actual, pred_cls, positive=1):
    tp = int(((pred_cls == positive) & (actual == positive)).sum())
    fn = int(((pred_cls != positive) & (actual == positive)).sum())
    fp = int(((pred_cls == positive) & (actual != positive)).sum())
    tn = int(((pred_cls != positive) & (actual != positive)).sum())
    return {"TP": tp, "FN": fn, "FP": fp, "TN": tn}


def check_run(ows: Path, workdir: Path) -> dict:
    """수강생 PC처럼 .ows와 CSV 두 개만 있는 폴더에서 연다 → 위젯 상태와 출력을 확인하고 리허설 값을 모은다."""
    from Orange.evaluation import AUC, CA, F1, Precision, Recall

    scheme, w = ow.run_ows(ows, [CK / TRAIN, CK / SCORE], workdir / "D3run")
    lv = lambda title, out: ow.link_value(scheme, title, out)  # noqa: E731
    problems: list[str] = []

    data = lv(T["file"], "data")
    dom = ow.domain_summary(data)
    if dom["class"] != ["late_30d"] or not {"buyer_id", "ref_date", "sample_weight", "split_suggested"} <= set(dom["metas"]):
        problems.append(f"File 역할이 복원되지 않음: {dom}")
    if {"fx_krw_close", "usd_krw_close"} & {v.name for v in data.domain.attributes}:
        problems.append("fx_krw_close · usd_krw_close가 특성으로 남음(skip이어야 함)")
    if len(w[T["sel"]].conditions) != 1:
        problems.append(f"Select Rows 조건 {len(w[T['sel']].conditions)}개(1개여야 함)")
    sample, remain = lv(T["ds1"], "data_sample"), lv(T["ds1"], "remaining_data")
    pos, over, train = lv(T["sel"], "matching_data"), lv(T["ds2"], "data_sample"), lv(T["concat"], "data")
    ts = w[T["ts"]]
    if ts.resampling != 5:
        problems.append("Test and Score가 Test on test data가 아님")
    if ts.class_selection != "1":
        problems.append(f"Test and Score Target class = {ts.class_selection!r}")
    res = lv(T["ts"], "evaluations_results")
    names = list(res.learner_names)
    metrics = {}
    for i, nm in enumerate(names):
        metrics[nm] = {
            "AUC": round(float(AUC(res, target=1)[i]), 4),
            "CA": round(float(CA(res)[i]), 4),
            "F1": round(float(F1(res, target=1)[i]), 4),
            "Precision": round(float(Precision(res, target=1)[i]), 4),
            "Recall": round(float(Recall(res, target=1)[i]), 4),
        }
    actual = res.actual.astype(int)
    gb_idx = names.index("Gradient Boosting") if "Gradient Boosting" in names else 0
    cm_gb = confusion(actual, res.predicted[gb_idx].astype(int))
    roc, cal = w[T["roc"]], w[T["cal"]]
    if (roc.fp_cost, roc.fn_cost, roc.target_index) != (1, 5, 1):
        problems.append(f"ROC 설정 {(roc.fp_cost, roc.fn_cost, roc.target_index)}")
    if (cal.target_index, list(cal.selected_classifiers), cal.output_calibration, cal.score) != (1, [0], 0, 0):
        problems.append(f"Calibration 설정 {(cal.target_index, cal.selected_classifiers, cal.output_calibration, cal.score)}")
    model = lv(T["cal"], "calibrated_model")
    if model is None:
        problems.append("Calibration Plot이 모델을 내보내지 않음")
    gbw = w[T["gb"]]
    gb_method = gbw.editor.learner_class.__name__ if gbw.editor.learner_class else None
    xgb = gbw.xgb_editor
    if gb_method != "XGBLearner" or (xgb.n_estimators, round(float(xgb.learning_rate), 3), xgb.max_depth) != (100, 0.1, 3):
        problems.append(f"Gradient Boosting 설정 {gb_method} {(xgb.n_estimators, xgb.learning_rate, xgb.max_depth)}")

    # Predictions: 150행 확률 · 시험셋 확률
    preds = lv(T["pred"], "annotated")
    preds_t = lv(T["pred_t"], "annotated")
    p_cols = [m.name for m in preds.domain.metas if m.name.endswith("(1)")]
    p150 = preds.get_column(p_cols[0]).astype(float) if p_cols else np.array([])
    pt_cols = [m.name for m in preds_t.domain.metas if m.name.endswith("(1)")]
    ptest = preds_t.get_column(pt_cols[0]).astype(float) if pt_cols else np.array([])
    ytest = preds_t.get_column("late_30d").astype(int) if preds_t is not None else np.array([])

    # 비용표(사이트 Step 4 🔵: t = 0.10 · 0.17 · 0.25 · 0.35 · 0.50, 비용 = 5·FN + FP) — 시험셋 확률 기준
    cost_rows = []
    for t in (0.10, 0.17, 0.25, 0.35, 0.50):
        fn = int(((ptest < t) & (ytest == 1)).sum())
        fp = int(((ptest >= t) & (ytest == 0)).sum())
        cost_rows.append({"t": t, "FN": fn, "FP": fp, "cost": 5 * fn + fp})

    def band(p, t):
        return np.where(p >= t, "C", np.where(p >= t / 2, "B", np.where(p >= t / 4, "A", "S")))

    grades = {}
    for t in (1 / 6, 0.17):
        g = band(p150, t)
        grades[f"{t:.4f}"] = {k: int((g == k).sum()) for k in "SABC"}

    # Save Data: 수강생이 Save를 누른 것처럼 저장해 열 순서 확인
    saved = {}
    for title in (T["save"], T["save_t"]):
        sw = w[title]
        try:
            sw.do_save()
            fn = Path(sw.filename)
            if fn.is_file():
                import openpyxl

                ws = openpyxl.load_workbook(fn, read_only=True).worksheets[0]
                head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
                saved[title] = {"file": fn.name, "rows": ws.max_row - 1, "columns_first5": head[:5],
                                "columns_last3": head[-3:], "n_columns": len(head)}
        except Exception as e:  # noqa: BLE001
            problems.append(f"{title} 저장 실패: {e!r}")

    # Explain Model 상위 10(평균 |SHAP|) — xgboost ≥ 3.3은 sklearn 기본값 enable_categorical=True라
    # shap(0.50+)이 '범주 분할'로 보고 거부한다(Orange는 이미 원-핫으로 바꿔 넣는데도). 환경 문제라 경고로 남긴다.
    xm = w[T["xm"]]
    top10 = []
    env_warnings: list[str] = []
    r = xm.results
    if r is not None:
        vals = np.abs(r.x[xm.target_index]).mean(axis=0)
        order = np.argsort(-vals)[:10]
        top10 = [{"feature": r.names[i], "mean_abs_shap": round(float(vals[i]), 4)} for i in order]
    else:
        msg = " ".join(m.formatted for m in xm.Error.active) if hasattr(xm.Error, "active") else ""
        if "Categorical split" in msg:
            import xgboost

            env_warnings.append(
                f"Explain Model 실패(xgboost {xgboost.__version__}): xgboost 3.3부터 enable_categorical 기본값이 True라 "
                "shap이 거부한다 → 포터블 Orange의 xgboost를 3.2.0 이하로 고정하거나, 그 자리에서 Gradient Boosting "
                "Method를 'Gradient Boosting (scikit-learn)'(같은 100·0.1·3)으로 바꿔 Step 7을 진행")
        else:
            problems.append(f"Explain Model 결과 없음: {msg or '원인 미상'}")
    if xm.target_index != 1:
        problems.append(f"Explain Model Target = {xm.target_index}")

    # Explain Prediction: 수강생처럼 Data Table(예측)에서 확률이 가장 높은 바이어 1행을 고른다
    local = {}
    tp, xp = w[T["t_pred"]], w[T["xp"]]
    i_top = int(np.argmax(p150))
    ncol = len(preds.domain.variables) + len(preds.domain.metas)
    tp.set_selection([i_top], list(range(ncol)))
    tp.commit.now()
    ow.settle(scheme, timeout=600)
    xr = getattr(xp, "_OWExplainPrediction__results", None)
    if xp.target_index != 1:
        problems.append(f"Explain Prediction Target = {xp.target_index}")
    if xr is not None:
        contrib = np.asarray(xr.values[xp.target_index])[0]
        names_x = [a.name for a in xr.transformed_data.domain.attributes]
        order = np.argsort(-np.abs(contrib))[:5]
        local = {"buyer_id": str(preds[i_top]["buyer_id"]), "pd_30d_calibrated": round(float(p150[i_top]), 4),
                 "model_output": round(float(np.asarray(xr.predictions)[0][xp.target_index]), 4),
                 "base_value": round(float(np.asarray(xr.base_value)[xp.target_index]), 4),
                 "top5": [{"feature": names_x[j], "contribution": round(float(contrib[j]), 4)} for j in order]}
    elif not env_warnings:
        problems.append("Explain Prediction 결과 없음")

    out = {
        "file_roles": dom,
        "split": {"train_rows": len(sample), "train_pos": int((sample.Y == 1).sum()), "test_rows": len(remain),
                  "test_pos": int((remain.Y == 1).sum())},
        "oversample": {"select_rows_pos": len(pos), "data_sampler2": len(over), "train_after": len(train),
                       "train_pos_after": int((train.Y == 1).sum())},
        "test_and_score": metrics,
        "confusion_gb": cm_gb,
        "roc": {"fp_cost": roc.fp_cost, "fn_cost": roc.fn_cost, "target_prior_pct": float(roc.target_prior)},
        "calibration_output": type(model).__name__ if model is not None else None,
        "gb_method": gb_method,
        "predictions": {"rows": len(preds), "prob_column": p_cols[0] if p_cols else None,
                        "lt_0_01": int((p150 < 0.01).sum()), "min": round(float(p150.min()), 4),
                        "median": round(float(np.median(p150)), 4), "max": round(float(p150.max()), 4),
                        "mean": round(float(p150.mean()), 4)},
        "test_predictions": {"rows": len(ptest), "pos": int(ytest.sum())},
        "cost_table_test": cost_rows,
        "grades_150_by_t": grades,
        "saved_files": saved,
        "explain_model_top10": top10,
        "explain_prediction_top_buyer": local,
    }
    import xgboost

    out["xgboost"] = xgboost.__version__
    gb_idx_cal = names.index("Gradient Boosting") if "Gradient Boosting" in names else 0
    raw = {"learner": lv(T["gb"], "learner"), "data": data, "results": res, "gb_idx": gb_idx_cal,
           "ptest_cal": ptest, "ytest": ytest, "p150": p150}
    return {"problems": problems, "warnings": env_warnings, "values": out, "raw": raw}


def _roles(table):
    """File 위젯과 같은 역할로 도메인을 다시 만든다(late_30d = target · 식별·보조 열 = meta · 환율 수준 = 뺌)."""
    from Orange.data import Domain

    dom = table.domain
    allv = list(dom.attributes) + list(dom.class_vars) + list(dom.metas)
    by = {v.name: v for v in allv}
    metas = [by[n] for n in ("buyer_id", "ref_date", "sample_weight", "split_suggested") if n in by]
    skip = {"late_30d", "buyer_id", "ref_date", "sample_weight", "split_suggested", "fx_krw_close", "usd_krw_close"}
    attrs = [v for v in allv if v.name not in skip and v.is_primitive()]
    return table.transform(Domain(attrs, by["late_30d"], metas))


def _chain_auc(learner, train, test, seed: int = 42) -> dict:
    """사이트 체인 그대로: 학습셋 연체만 ×2 복원 추출로 붙이고(오버샘플) Test on test data."""
    from Orange.data import Table
    from Orange.evaluation import AUC, TestOnTestData
    from sklearn.metrics import average_precision_score

    pos = np.flatnonzero(train.Y == 1)
    rng = np.random.RandomState(seed)
    extra = train[rng.choice(pos, size=2 * len(pos), replace=True)]
    train_os = Table.concatenate([train, extra])
    res = TestOnTestData(store_data=True)(train_os, test, [learner])
    p1 = res.probabilities[0][:, 1]
    return {"auc": round(float(AUC(res, target=1)[0]), 4), "pr_auc": round(float(average_precision_score(res.actual, p1)), 4),
            "n_train": len(train), "n_train_os": len(train_os), "n_test": len(test), "pos_test": int((test.Y == 1).sum())}


def rehearsal(vals: dict, raw: dict, v: dict) -> dict:
    """D-1 리허설 대용 기대값 — tools/train_day3_model.py 보고서 §6이 읽는 키 + 자세한 값."""
    from Orange.data import Table

    res, gi = raw["results"], raw["gb_idx"]
    p_raw = res.probabilities[gi][:, 1]
    pt, yt = raw["ptest_cal"], raw["ytest"]
    grid = np.round(np.arange(0.01, 1.0, 0.01), 2)
    cost = np.array([5 * int(((pt < t) & (yt == 1)).sum()) + int(((pt >= t) & (yt == 0)).sum()) for t in grid])
    best = grid[cost == cost.min()]
    t_star = float(best[np.argmin(np.abs(best - 1 / 6))])          # 최소 비용 구간 안에서 이론값 1/6에 가장 가까운 점
    fn = int(((pt < t_star) & (yt == 1)).sum())
    fp = int(((pt >= t_star) & (yt == 0)).sum())
    p150 = raw["p150"]
    g = np.where(p150 >= t_star, "C", np.where(p150 >= t_star / 2, "B", np.where(p150 >= t_star / 4, "A", "S")))
    sp, m = vals["split"], vals["test_and_score"]["Gradient Boosting"]
    data = raw["data"]
    def labels(table, name):   # 범주형 meta는 값 번호로 나오므로 글자로 바꾼다
        var = table.domain[name]
        return np.array([var.str_val(x) for x in table.get_column(name)])

    split = labels(data, "split_suggested")
    group = _chain_auc(raw["learner"], data[split == "train"], data[split == "test"])
    panel_t = _roles(Table.from_file(str(CK / "d3_features_panel.csv")))
    psplit = labels(panel_t, "split_suggested")
    panel = _chain_auc(raw["learner"], panel_t[psplit == "train"], panel_t[psplit == "test"])
    return {
        "orange_version": v.get("Orange3"), "xgboost_version": v.get("xgboost"),
        "basic": {"n_train": sp["train_rows"], "pos_train": sp["train_pos"], "n_train_os": vals["oversample"]["train_after"],
                  "n_test": sp["test_rows"], "pos_test": sp["test_pos"], "auc": m["AUC"], "ca": m["CA"], "f1": m["F1"],
                  "precision": m["Precision"], "recall": m["Recall"], "mean_p": round(float(p_raw.mean()), 4),
                  "t_cost_calibrated": t_star, "t_cost_min_range": [float(best.min()), float(best.max())],
                  "cost_at_t": int(cost.min()), "fn_cal": fn, "fp_cal": fp},
        "group": {"auc": group["auc"], "pr_auc": group["pr_auc"], "pos_test": group["pos_test"],
                  "n_train": group["n_train"], "n_test": group["n_test"]},
        "panel": {"auc": panel["auc"], "pr_auc": panel["pr_auc"], "n_train": panel["n_train"], "n_test": panel["n_test"],
                  "pos_test": panel["pos_test"]},
        "predictions": {"mean_p": vals["predictions"]["mean"], "t": t_star, "grades": {k: int((g == k).sum()) for k in "SABC"}},
        "explain_top5": [d["feature"] for d in vals["explain_model_top10"][:5] if isinstance(d, dict)],
    }


REPORT_PLACEHOLDER = "- (Orange 기대값 파일 없음)"


def fill_report(o: dict) -> bool:
    """d3_model_report.md §6의 자리표시 한 줄만 리허설 값으로 바꾼다(모델·체크포인트는 다시 만들지 않는다)."""
    rp = ANSWERS / "d3_model_report.md"
    if not rp.is_file():
        return False
    lines = rp.read_text(encoding="utf-8").splitlines()
    idx = [i for i, ln in enumerate(lines) if ln.startswith(REPORT_PLACEHOLDER) or ln.startswith("- 실행 환경: Orange")]
    if not idx:
        return False
    b = o["basic"]
    txt = [
        f"- 실행 환경: Orange {o['orange_version']} · xgboost {o['xgboost_version']} (`tools/build_day3_orange.py --expected`가 "
        f"`labs/day3/d3_orange_model.ows`를 수강생 PC처럼 다시 열어 실행 — D-1 리허설 대용)",
        f"- 🟢 300행 · Data Sampler 70% 층화·Replicable → 학습 {b['n_train']}행(양성 {b['pos_train']}, 오버샘플 후 {b['n_train_os']}행) "
        f"/ 테스트 {b['n_test']}행(양성 {b['pos_test']})",
        f"  - Test and Score(Test on test data): **AUC {b['auc']} · CA {b['ca']} · F1 {b['f1']} · Precision {b['precision']} · "
        f"Recall {b['recall']}**(임계값 0.5)",
        f"  - 확률 평균 {b['mean_p']}(오버샘플로 부풀어 있음) → Sigmoid 보정 후 비용 최적 임계값 {b['t_cost_calibrated']}"
        f"(비용 최소 구간 {b['t_cost_min_range'][0]}–{b['t_cost_min_range'][1]}) · 이때 FN {b['fn_cal']} · FP {b['fp_cal']}",
        f"- 🔵 바이어 단위 분할(split_suggested, 같은 체인): AUC {o['group']['auc']} (테스트 양성 {o['group']['pos_test']})",
        f"- 🔵 패널 시간 분할(split_suggested, 같은 체인): AUC {o['panel']['auc']} · PR-AUC {o['panel']['pr_auc']} "
        f"(학습 {o['panel']['n_train']:,} / 검증 {o['panel']['n_test']:,}행)",
        f"- 🟢 Predictions(2026-09-30, 150개사, 보정 모델): pd 평균 {o['predictions']['mean_p']} · t = {o['predictions']['t']} "
        f"기준 등급 {o['predictions']['grades']}",
        f"- Explain Model 상위 5: {', '.join(o['explain_top5']) or '계산 안 됨(xgboost ≥ 3.3 — orange_expected.json warnings)'}",
    ]
    i0 = idx[0]
    i1 = i0 + 1
    while i1 < len(lines) and lines[i1].startswith(("- 🟢", "  - ", "- 🔵", "- Explain Model 상위")):
        i1 += 1
    lines[i0:i1] = txt
    rp.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--expected", action="store_true", help="강사 리허설 기대값을 ../instructor/day3/answers/에 쓴다")
    ap.add_argument("--keep", action="store_true", help="임시 폴더를 지우지 않는다(디버그)")
    args = ap.parse_args()
    v = ow.check_versions()
    tmp = Path(tempfile.mkdtemp(prefix="d3_orange_"))
    try:
        path, seen = build(tmp)
        print(f"생성: {path.relative_to(REPO)} ({path.stat().st_size:,} bytes) · 학습 연체 {seen['train_pos']} → "
              f"복제 {seen['n_over']}")
        info = ow.load_check(path)
        print(f"다시 열기: 노드 {info['nodes']} · 링크 {info['links']} · 메모 {info['annotations']} · 알 수 없는 위젯 0")
        chk = check_run(path, tmp)
        vals = chk["values"]
        print("실행: 분할", vals["split"], "· 오버샘플", vals["oversample"])
        for nm, m in vals["test_and_score"].items():
            print(f"  {nm}: {m}")
        print("  Confusion(GB)", vals["confusion_gb"], "· 보정 출력", vals["calibration_output"],
              "· 예측", vals["predictions"])
        print("  저장 파일", vals["saved_files"])
        print("  비용표(시험셋)", vals["cost_table_test"])
        print("  등급(150곳, t별)", vals["grades_150_by_t"])
        print("  Explain Model 상위 10", [d["feature"] if isinstance(d, dict) else d for d in vals["explain_model_top10"]])
        print("  Explain Prediction", vals["explain_prediction_top_buyer"])
        for msg in chk["warnings"]:
            print("[환경 경고]", msg)
        if chk["problems"]:
            print("문제:", *chk["problems"], sep="\n  - ")
            return 1
        print("검사 통과 — 설정 복원 · 신호 전달 · 저장까지 문제 없음")
        if args.expected:
            ANSWERS.mkdir(parents=True, exist_ok=True)
            reh = rehearsal(vals, chk["raw"], v)
            doc = {"generated": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                   "source": "tools/build_day3_orange.py --expected (Orange 헤드리스 실행 = D-1 리허설 대용)",
                   **reh, "versions": v, "workflow": str(path.relative_to(REPO)), "nodes": info["titles"],
                   "warnings": chk["warnings"], "detail": vals, "notes": NOTES}
            dst = ANSWERS / "orange_expected.json"
            dst.write_text(json.dumps(doc, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            print(f"강사 기대값: {dst}")
            print("  basic", reh["basic"], "· group", reh["group"], "· panel", reh["panel"], "· 예측", reh["predictions"])
            if fill_report(reh):
                print("  d3_model_report.md §6 자리표시를 채웠다(모델·체크포인트는 그대로)")
        return 0
    finally:
        if not args.keep:
            shutil.rmtree(tmp, ignore_errors=True)
        else:
            print("임시 폴더:", tmp)


if __name__ == "__main__":
    sys.exit(main())
